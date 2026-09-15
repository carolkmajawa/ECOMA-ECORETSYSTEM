# apps/reports/services/excel_generator.py

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime


class ExcelReportGenerator:
    """Generate Excel reports"""

    HEADER_FILL = PatternFill(start_color='0B6B80', end_color='0B6B80', fill_type='solid')
    HEADER_FONT = Font(color='FFFFFF', bold=True)

    def __init__(self, report_data, group_name):
        self.data = report_data
        self.group_name = group_name
        self.wb = Workbook()
        self.wb.remove(self.wb.active)  # drop the default blank sheet

    def generate(self, output_filename):
        self._add_summary_sheet()
        self._add_members_sheet()
        self._add_transactions_sheet()
        self.wb.save(output_filename)
        return output_filename

    def _style_header_row(self, ws, row_num, num_cols):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal='center')

    def _add_summary_sheet(self):
        ws = self.wb.create_sheet('Summary')
        ws['A1'] = f"{self.group_name} - Cycle Report"
        ws['A1'].font = Font(size=16, bold=True, color='C50982')
        ws['A2'] = f"Generated: {datetime.now().strftime('%B %d, %Y at %H:%M')}"

        group = self.data.get('group_info', {}) if isinstance(self.data.get('group_info'), dict) else {}
        contrib = self.data.get('contributions', {}) if isinstance(self.data.get('contributions'), dict) else {}
        loans = self.data.get('loans', {}) if isinstance(self.data.get('loans'), dict) else {}
        financial = self.data.get('financial_summary', {}) if isinstance(self.data.get('financial_summary'), dict) else {}

        rows = [
            ('Group Name', group.get('name', 'N/A')),
            ('Group Code', group.get('code', 'N/A')),
            ('Total Members', group.get('total_members', 0)),
            ('Total Savings', contrib.get('total_savings', 0)),
            ('Total Loans', contrib.get('total_loans', 0)),
            ('Total Loan Count', loans.get('total_loans', 0)),
            ('Active Loans', loans.get('active_loans', 0)),
            ('Total Contributions', financial.get('total_contributions', 0)),
        ]

        row_num = 4
        for label, value in rows:
            ws.cell(row=row_num, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row_num, column=2, value=value)
            row_num += 1

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 25

    def _add_members_sheet(self):
        members = self.data.get('members', [])
        if not isinstance(members, list):
            members = []

        ws = self.wb.create_sheet('Members')
        headers = ['#', 'Full Name', 'Phone', 'Gender', 'Address', 'Status']
        ws.append(headers)
        self._style_header_row(ws, 1, len(headers))

        for idx, member in enumerate(members, 1):
            if not isinstance(member, dict):
                continue
            ws.append([
                idx,
                member.get('full_name', 'N/A'),
                member.get('phone_number', 'N/A'),
                member.get('gender', 'N/A'),
                member.get('address', 'N/A'),
                'Active' if member.get('is_active') else 'Inactive',
            ])

        for col, width in zip('ABCDEF', [5, 25, 15, 10, 25, 10]):
            ws.column_dimensions[col].width = width

    def _add_transactions_sheet(self):
        transactions = self.data.get('transactions', {})
        if not isinstance(transactions, dict):
            transactions = {}
        txn_list = transactions.get('transactions', [])

        ws = self.wb.create_sheet('Transactions')
        headers = ['Date', 'Type', 'Description', 'Member', 'Amount', 'Recorded By']
        ws.append(headers)
        self._style_header_row(ws, 1, len(headers))

        for t in txn_list:
            ws.append([
                t.get('date', 'N/A'),
                t.get('transaction_type', 'N/A'),
                t.get('description', 'N/A'),
                t.get('member', 'N/A'),
                t.get('amount', 0),
                t.get('recorded_by', 'N/A'),
            ])

        for col, width in zip('ABCDEF', [20, 15, 30, 20, 12, 20]):
            ws.column_dimensions[col].width = width